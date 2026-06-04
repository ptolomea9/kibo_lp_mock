from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook()
ws = wb.active
ws.title = "PPC OMS vs Organic OMS"

navy = "0D1B3E"
green_bg = "E8F5E9"
blue_bg = "E3F2FD"
header_font = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
header_fill = PatternFill("solid", fgColor=navy)
body_font = Font(size=10, name="Calibri")
green_font = Font(size=10, name="Calibri", color="1B5E20")
blue_font = Font(size=10, name="Calibri", color="1565C0")
bold_green = Font(size=10, name="Calibri", color="1B5E20", bold=True)
bold_blue = Font(size=10, name="Calibri", color="1565C0", bold=True)
cat_font = Font(bold=True, size=10, name="Calibri", color=navy)
cat_fill = PatternFill("solid", fgColor="F0F0F0")
thin = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)
wrap = Alignment(wrap_text=True, vertical="top")

ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 46
ws.column_dimensions["C"].width = 46
ws.column_dimensions["D"].width = 16

# Title
ws.merge_cells("A1:D1")
ws["A1"] = "Kibo OMS: PPC Landing Page vs. Existing Organic Page"
ws["A1"].font = Font(bold=True, size=14, name="Calibri", color=navy)
ws.merge_cells("A2:D2")
ws["A2"] = "Why a dedicated PPC page will improve Quality Score LP Experience from BELOW_AVERAGE"
ws["A2"].font = Font(size=10, name="Calibri", color="666666")

headers = [
    "Factor",
    "Existing OMS Page\n(kibocommerce.com/platform/order-management/)",
    "PPC OMS Page\n(proposed /ppc/oms/)",
    "Verdict",
]
for col, val in enumerate(headers, 1):
    c = ws.cell(row=4, column=col, value=val)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    c.border = thin
ws.row_dimensions[4].height = 40

data = [
    ("CAT", "PAGE WEIGHT & PERFORMANCE"),
    ("ROW", "HTML Size", "457 KB (4.5x over 100 KB benchmark)", "~35 KB (well under benchmark)", "PPC"),
    ("ROW", "Total Scripts", "76 scripts from 20 external domains", "0 external (6 in production: GTM, HS form, GAds, LI, GCLID, li_fat_id)", "PPC"),
    ("ROW", "Stylesheets", "48 stylesheets", "1 (inline CSS)", "PPC"),
    ("ROW", "Images", "103 images", "8 customer logos (lazy-loaded)", "PPC"),
    ("ROW", "Console Errors", "SyntaxError on load + HubSpot widget warning", "None", "PPC"),
    ("CAT", "CONVERSION ARCHITECTURE"),
    ("ROW", "Forms on Page", "ZERO. Both CTAs link off-page to /request-a-demo/", "Embedded form above the fold + repeated at bottom CTA", "PPC"),
    ("ROW", "Primary CTA", '"Watch the Demo" links to gated form (misleading)', '"Request a Demo" with inline form (honest CTA)', "PPC"),
    ("ROW", "Analyst Proof", "Forrester Wave buried at bottom in resource card", "Forrester badge above fold in analyst bar + hero", "PPC"),
    ("ROW", "Social Proof", "45 logos, zero named quotes, zero metrics", "Attributed testimonial (20+ DCs, 400 stores)", "PPC"),
    ("ROW", "Tier CTAs", "Three tiers with features but no CTA buttons", "Starter/Essentials/Advanced with per-tier CTAs", "PPC"),
    ("CAT", "KEYWORD-CONTENT ALIGNMENT"),
    ("ROW", "H1 Keyword Match", '"Order Management" (no "system", no "enterprise")', '"Enterprise Order Management System" (full match)', "PPC"),
    ("ROW", "Meta Description", 'Brand-speak: "Promise accurately and fulfill anywhere"', 'Search-intent: "unifies inventory, routing, fulfillment... Request a demo"', "PPC"),
    ("ROW", '"Ecommerce" Usage', 'Barely appears. Uses "commerce" (Kibo branding)', "15+ natural instances throughout visible copy", "PPC"),
    ("ROW", '"Software" Usage', "Absent from visible copy", "In section headings and tier descriptions", "PPC"),
    ("ROW", '"Enterprise" Usage', "Only in title tag, not in visible body", "In hero, section copy, analyst bar, tier descriptions", "PPC"),
    ("CAT", "CONTENT QUALITY"),
    ("ROW", "Heading Typos", '"How it KIBO OMS Works" and "Packaging Packaging Tiers"', "All clean, proofread", "PPC"),
    ("ROW", "Duplicate Content", "Two sections share identical body paragraphs", "All unique copy per section", "PPC"),
    ("ROW", "FAQ Section", "No FAQ section, no structured FAQ content", "6 FAQs in two-column expandable layout", "PPC"),
    ("ROW", "Schema Markup", "1 broken/empty schema block", "FAQPage (6 Q&As) + SoftwareApplication (3 tiers)", "PPC"),
    ("ROW", "Page Length", "12+ sections, much shallow/duplicated", "10 focused sections, each with substantive copy", "PPC"),
    ("CAT", "WHAT THE ORGANIC PAGE DOES BETTER"),
    ("ROW", "Brand Awareness", "Full product marketing narrative, Kibo story", "Conversion-focused, less brand storytelling", "ORG"),
    ("ROW", "Feature Depth", "45+ H4 feature labels, tabbed capability matrix", "Streamlined 6-card feature grid", "ORG"),
    ("ROW", "SEO Authority", "Indexed, ranks for branded queries, link equity", "noindex/nofollow (PPC only, by design)", "ORG"),
]

row = 5
for d in data:
    if d[0] == "CAT":
        c = ws.cell(row=row, column=1, value=d[1])
        c.font = cat_font
        c.fill = cat_fill
        c.border = thin
        for col in range(2, 5):
            ws.cell(row=row, column=col).fill = cat_fill
            ws.cell(row=row, column=col).border = thin
        ws.row_dimensions[row].height = 22
    else:
        _, factor, existing, ppc, verdict = d
        ws.cell(row=row, column=1, value=factor).font = Font(size=10, name="Calibri", bold=True)
        ws.cell(row=row, column=2, value=existing).font = body_font
        ws.cell(row=row, column=3, value=ppc).font = body_font
        v = "PPC wins" if verdict == "PPC" else "Organic wins"
        ws.cell(row=row, column=4, value=v).font = body_font
        for col in range(1, 5):
            ws.cell(row=row, column=col).alignment = wrap
            ws.cell(row=row, column=col).border = thin
        if verdict == "PPC":
            ws.cell(row=row, column=3).font = green_font
            ws.cell(row=row, column=3).fill = PatternFill("solid", fgColor=green_bg)
            ws.cell(row=row, column=4).font = bold_green
        else:
            ws.cell(row=row, column=2).font = blue_font
            ws.cell(row=row, column=2).fill = PatternFill("solid", fgColor=blue_bg)
            ws.cell(row=row, column=4).font = bold_blue
        ws.row_dimensions[row].height = 52
    row += 1

row += 1
ws.merge_cells(f"A{row}:D{row}")
ws.cell(row=row, column=1, value="Summary").font = Font(bold=True, size=12, name="Calibri", color=navy)
row += 1
ws.merge_cells(f"A{row}:D{row}")
c = ws.cell(
    row=row,
    column=1,
    value=(
        "The existing OMS page scores 4/10 overall as a paid search landing page. "
        "It was built as a product marketing page and repurposed for paid search. "
        "The 100% BELOW_AVERAGE LP Experience scores are driven by: "
        "457KB HTML with 76 scripts (page speed), brand language instead of search keywords "
        "(content relevance), and no on-page form (navigation ease). "
        "A dedicated PPC page controls all three factors without touching the organic page."
    ),
)
c.font = Font(size=10, name="Calibri", color="333333")
c.alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[row].height = 72

row += 2
ws.merge_cells(f"A{row}:D{row}")
c = ws.cell(
    row=row,
    column=1,
    value=(
        "Root Cause: Google LP Experience evaluates page speed, content-keyword relevance, "
        "and ease of navigation. The organic page fails all three. "
        "The fix is not incremental. A dedicated PPC LP is the fastest path to AVERAGE or ABOVE_AVERAGE."
    ),
)
c.font = Font(size=10, name="Calibri", color="666666", italic=True)
c.alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[row].height = 48

wb.save("PPC_OMS_vs_Organic_Comparison.xlsx")
print("Saved PPC_OMS_vs_Organic_Comparison.xlsx")
